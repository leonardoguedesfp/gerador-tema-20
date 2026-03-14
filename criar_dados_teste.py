#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Cria arquivos de teste para o Gerador de Petições.
Gera: planilha de dados e 4 modelos Word com variáveis e bloco condicional.
"""

import os
from pathlib import Path
from docx import Document
from docx.shared import Pt
import openpyxl


PASTA_ENTRADA = Path("entrada")
PASTA_ENTRADA.mkdir(exist_ok=True)


def criar_modelo(caminho: str, genero: str, qtd_rts: int):
    """Cria um modelo Word com variáveis e bloco condicional."""
    doc = Document()

    # Configurar estilo padrão
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)

    pronome = "o" if genero == "M" else "a"
    artigo = "O" if genero == "M" else "A"
    reclamante = f"{artigo} reclamante"

    # Cabeçalho
    doc.add_paragraph(
        "EXCELENTÍSSIMO(A) SENHOR(A) JUIZ(A) DO TRABALHO DA ___ VARA DO "
        "TRABALHO DE BRASÍLIA/DF"
    )
    doc.add_paragraph("")

    # Qualificação
    doc.add_paragraph(
        "{Nome}, {QualificacaoAutor}, "
        f"vem, respeitosamente, à presença de Vossa Excelência, por "
        f"intermédio d{pronome}s advogad{pronome}s abaixo assinad{pronome}s, "
        f"propor a presente"
    )
    doc.add_paragraph("")

    doc.add_paragraph("RECLAMAÇÃO TRABALHISTA")
    doc.add_paragraph("")

    doc.add_paragraph(
        "em face de CAIXA DE PREVIDÊNCIA DOS FUNCIONÁRIOS DO BANCO DO BRASIL "
        "— PREVI, pessoa jurídica de direito privado..."
    )
    doc.add_paragraph("")

    # Dados do desligamento
    doc.add_paragraph(
        f"{reclamante} foi desligad{pronome} do Banco do Brasil em "
        "{DataDesligamento}, percebendo benefício de complementação de "
        "aposentadoria junto à PREVI no valor de {ValorBeneficio}."
    )
    doc.add_paragraph("")

    # RTs
    if qtd_rts == 2:
        doc.add_paragraph(
            f"{reclamante} ajuizou ações trabalhistas, sob os números "
            "{RT_Verbas_1} e {RT_Verbas_2}, nas quais foram reconhecidas "
            "verbas trabalhistas que deveriam ter sido incorporadas ao "
            "cálculo do benefício."
        )
    else:
        doc.add_paragraph(
            f"{reclamante} ajuizou ação trabalhista, sob o número "
            "{RT_Verbas_1}, na qual foram reconhecidas verbas trabalhistas "
            "que deveriam ter sido incorporadas ao cálculo do benefício."
        )
    doc.add_paragraph("")

    # RT Indenizatória
    doc.add_paragraph(
        f"{reclamante} já ajuizou anteriormente ação indenizatória "
        "(processo nº {RT_Indenizatoria}), tendo sido extinta sem resolução "
        "do mérito."
    )
    doc.add_paragraph("")

    # Tópico 2 — Prescrição
    doc.add_paragraph("2. DA PRESCRIÇÃO")
    doc.add_paragraph("{TrechoPrescricao}")
    doc.add_paragraph("")

    # Tópico 4.1
    doc.add_paragraph("4.1. DOS REFLEXOS NAS CONTRIBUIÇÕES")
    doc.add_paragraph(
        "As verbas reconhecidas judicialmente devem repercutir no cálculo "
        "do benefício de complementação de aposentadoria."
    )
    doc.add_paragraph("")

    # Bloco condicional 4.2
    doc.add_paragraph("{SE:PreservacaoSP}")
    doc.add_paragraph("4.2. DA PRESERVAÇÃO DO SALÁRIO DE PARTICIPAÇÃO")
    doc.add_paragraph(
        f"Registre-se que {pronome} salário de participação d{pronome} "
        f"reclamante foi preservado {'{'}PreservacaoSP{'}'}, conforme "
        "documentação anexa."
    )
    doc.add_paragraph("{FIM_SE:PreservacaoSP}")

    # Tópico 5
    doc.add_paragraph("5. DOS PEDIDOS")
    doc.add_paragraph(
        "Diante do exposto, requer a condenação da reclamada ao pagamento "
        "das diferenças de complementação de aposentadoria."
    )
    doc.add_paragraph("")

    doc.add_paragraph("Dá-se à causa o valor de R$ 65.000,00 (sessenta e cinco mil reais).")
    doc.add_paragraph("")

    doc.add_paragraph("Brasília, {DataPeticao}.")
    doc.add_paragraph("")

    doc.add_paragraph("_______________________________")
    doc.add_paragraph("Advogado — OAB/DF 00000")

    doc.save(str(caminho))


def criar_planilha(caminho: str):
    """Cria a planilha de dados de teste."""
    wb = openpyxl.Workbook()

    # Aba de dados
    ws = wb.active
    ws.title = "Dados Clientes"

    cabecalhos = [
        "Nome", "Genero", "QualificacaoAutor", "DataDesligamento",
        "ValorBeneficio", "RT_Verbas_1", "Verbas_1", "TransitoJulgado_1",
        "RT_Verbas_2", "Verbas_2", "TransitoJulgado_2",
        "RT_Indenizatoria", "PreservacaoSP", "QtdRTs",
        "Grupo", "DataReferenciaSTJ", "TrechoPrescricao",
    ]
    ws.append(cabecalhos)

    # Teste 1: Masculino, 2 RTs, com preservação
    ws.append([
        "RINALDO DA SILVA SOARES",
        "M",
        "brasileiro, casado, bancário aposentado, CPF 000.111.222-33, "
        "RG 1.234.567 SSP/DF, residente na QNM 10 Conjunto A Casa 1, "
        "Ceilândia/DF, CEP 72215-101",
        "16/7/2025",
        "R$ 10.233,57",
        "0000263-76.2015.5.10.0010",
        "horas extras",
        "15/3/2020",
        "0000765-62.2017.5.10.0004",
        "diferenças salariais",
        "22/8/2021",
        "0000237-50.2021.5.10.0016",
        "em 6/2017, no importe de R$9.729,16",
        2,
        "1.1",
        "16/08/2018",
        "No caso em tela, o trânsito em julgado da primeira ação "
        "trabalhista ocorreu em 15/03/2020, portanto dentro do prazo "
        "prescricional de 5 anos previsto no art. 7º, XXIX, da CF/88.",
    ])

    # Teste 2: Feminino, 1 RT, sem preservação
    ws.append([
        "CARLA BEZERRA MARTINS DE ANDRADE",
        "F",
        "brasileira, solteira, bancária aposentada, CPF 333.444.555-66, "
        "RG 9.876.543 SSP/DF, residente na SQS 308 Bloco A Apt 101, "
        "Brasília/DF, CEP 70352-010",
        "22/3/2024",
        "R$ 8.450,00",
        "0001876-92.2014.5.10.0002",
        "horas extras e gratificação",
        "10/6/2019",
        "",
        "",
        "",
        "0001234-56.2022.5.10.0008",
        "",
        1,
        "2.1",
        "11/12/2020",
        "Quanto à prescrição, o trânsito em julgado ocorreu em "
        "10/06/2019, sendo certo que a presente ação foi ajuizada "
        "dentro do prazo quinquenal.",
    ])

    # Teste 3: Masculino, 1 RT, com preservação
    ws.append([
        "MARCOS ANTONIO FERREIRA LIMA",
        "M",
        "brasileiro, divorciado, bancário aposentado, CPF 777.888.999-00, "
        "RG 5.555.666 SSP/GO, residente na Rua 10 Qd 5 Lt 12, "
        "Goiânia/GO, CEP 74000-000",
        "5/1/2023",
        "R$ 12.100,00",
        "0002345-11.2016.5.10.0005",
        "adicional noturno",
        "28/11/2020",
        "",
        "",
        "",
        "0003456-78.2023.5.10.0012",
        "em 3/2019, no importe de R$11.200,00",
        1,
        "1.2",
        "16/08/2018",
        "Em relação à prescrição, aplica-se o prazo quinquenal, "
        "contado do trânsito em julgado da ação trabalhista originária.",
    ])

    # Aba de legenda (ignorada pelo sistema)
    ws2 = wb.create_sheet("Legenda e Instruções")
    ws2.append(["Esta aba é apenas para referência do usuário."])
    ws2.append(["O sistema ignora esta aba completamente."])

    wb.save(str(caminho))


def main():
    print("Criando arquivos de teste...")

    # Criar modelos
    for (genero, qtd_rts), caminho in [
        (("M", 1), PASTA_ENTRADA / "modelo_masculino_1RT.docx"),
        (("M", 2), PASTA_ENTRADA / "modelo_masculino_2RT.docx"),
        (("F", 1), PASTA_ENTRADA / "modelo_feminino_1RT.docx"),
        (("F", 2), PASTA_ENTRADA / "modelo_feminino_2RT.docx"),
    ]:
        criar_modelo(str(caminho), genero, qtd_rts)
        print(f"  Modelo criado: {caminho}")

    # Criar planilha
    criar_planilha(str(PASTA_ENTRADA / "dados_clientes.xlsx"))
    print(f"  Planilha criada: {PASTA_ENTRADA / 'dados_clientes.xlsx'}")

    print("Arquivos de teste criados com sucesso!")
    print("\nAgora execute: python gerar_peticoes.py")


if __name__ == "__main__":
    main()
